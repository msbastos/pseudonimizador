# SPDX-License-Identifier: AGPL-3.0-only
# Termo adicional (AGPL-3.0, seção 7(b)): preservar os avisos de autoria; ver CREDITOS.md.
import _base  # noqa: F401
import zipfile

import docx
import fabrica as F
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from pseudonimizador.leitura import ler
from pseudonimizador.revisao import Revisao
from pseudonimizador.saida.docx_layout import gerar_docx_layout


def _contaminar(caminho):
    """Hyperlink com alvo mailto e tooltip, alt text (o modelo padrão do python-docx já traz miniatura e customXml)."""
    d = docx.Document(str(caminho))
    rid = d.part.relate_to("mailto:" + F.EMAIL, RT.HYPERLINK, is_external=True)
    corpo = d.element.body
    corpo.append(parse_xml(
        f'<w:p {nsdecls("w", "r", "wp")}><w:hyperlink r:id="{rid}" w:tooltip="{F.PESSOA_1}">'
        '<w:r><w:t>contato</w:t></w:r></w:hyperlink>'
        '<w:r><w:drawing><wp:inline><wp:extent cx="1" cy="1"/>'
        f'<wp:docPr id="1" name="Figura" descr="{F.PESSOA_2}" title="{F.PESSOA_1}"/>'
        '</wp:inline></w:drawing></w:r></w:p>'))
    sect = corpo.find(qn("w:sectPr"))
    if sect is not None:
        sect.addprevious(corpo[-1])
    d.save(str(caminho))


def _gerar():
    pasta = F.pasta_temporaria("saida_docx")
    origem = F.contrato_docx(pasta / "c.docx")
    _contaminar(origem)
    with zipfile.ZipFile(origem) as z:
        assert any("thumbnail" in n for n in z.namelist()) and any(n.startswith("customXml/") for n in z.namelist())
    doc = ler(origem)
    res = Revisao(doc).fechar()
    destino = pasta / "c - pseudonimizado.docx"
    gerar_docx_layout(doc, res.trocas, destino, identificador="pseudonimizador:teste")
    return destino, res


def test_nada_em_claro_em_nenhuma_parte_do_arquivo():
    destino, _ = _gerar()
    with zipfile.ZipFile(destino) as z:
        bruto = " ".join(z.read(n).decode("utf-8", "replace") for n in z.namelist() if n.endswith((".xml", ".rels")))
    for valor in (F.PESSOA_1, F.PESSOA_2, F.CPF, F.EMAIL, "Autor Original", "Editor Sintético",
                  "Comentarista Sintético", "Conferir o CPF", "Revisor"):
        assert valor not in bruto, valor


def test_sem_miniatura_nem_customxml():
    destino, _ = _gerar()
    with zipfile.ZipFile(destino) as z:
        nomes = z.namelist()
    assert not [n for n in nomes if n.startswith("customXml/") or "thumbnail" in n], nomes


def test_layout_preservado_e_tokens_no_lugar():
    destino, res = _gerar()
    d = ler(destino)
    assert "LOCADOR: [[PESSOA_1]], brasileira" in d.texto
    assert "Cláusula 1ª. O aluguel mensal" in d.texto
    original = docx.Document(str(destino))
    p = original.paragraphs[1]
    assert p.runs[0].text == "LOCADOR: [[PESSOA_1]]" and p.runs[1].text == "" and p.runs[1].bold
    assert original.core_properties.identifier == "pseudonimizador:teste"


if __name__ == "__main__":
    from _runner import rodar
    rodar(globals())
