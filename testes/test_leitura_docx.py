# SPDX-License-Identifier: AGPL-3.0-only
# Termo adicional (AGPL-3.0, seção 7(b)): preservar os avisos de autoria; ver CREDITOS.md.
import _base  # noqa: F401
import fabrica as F
from pseudonimizador.leitura.docx_ import ler_docx
from pseudonimizador.leitura.modelo import OrigemDocx


def _doc():
    return ler_docx(F.contrato_docx(F.pasta_temporaria("leitura_docx") / "c.docx"))


def test_texto_inclui_corpo_tabela_cabecalho_hyperlink_insercao_e_caixa():
    d = _doc()
    assert d.tipo == "docx"
    assert "LOCADOR: Maria Aparecida Pereira, brasileira" in d.texto       # nome partido entre runs
    assert d.texto.count(F.PESSOA_1) >= 5   # parágrafo, assinatura, tabela, cabeçalho, caixa de texto
    assert F.EMAIL in d.texto                                               # hyperlink
    assert "Testemunha: " + F.PESSOA_2 in d.texto                           # revisão aceita
    assert "Fiadora: " + F.PESSOA_1 in d.texto                              # caixa de texto


def test_origens_alinhadas_ao_texto():
    d = _doc()
    assert len(d.origens) == len(d.texto)
    i = d.texto.index("recida Pereira")
    o = d.origens[i]
    assert isinstance(o, OrigemDocx) and o.run == 1 and o.pos == 0
    assert all(d.origens[k] is None for k, c in enumerate(d.texto) if c == "\n")


def _espera_erro(caminho):
    from pseudonimizador.leitura.modelo import ErroDeLeitura
    try:
        ler_docx(caminho)
    except ErroDeLeitura:
        return
    raise AssertionError("devia levantar ErroDeLeitura")


def test_texto_puro_ole_e_inexistente_viram_erro_de_leitura():
    pasta = F.pasta_temporaria("leitura_docx_erro")
    (pasta / "texto.docx").write_text("isto não é um docx", encoding="utf-8")
    (pasta / "ole.docx").write_bytes(bytes.fromhex("D0CF11E0A1B11AE1") + bytes(600))
    for nome in ("texto.docx", "ole.docx", "nao_existe.docx"):
        _espera_erro(pasta / nome)


def test_runs_em_sdt_inline_sao_lidos():
    import docx
    from docx.oxml import parse_xml
    from pseudonimizador.leitura.modelo import OrigemDocx
    pasta = F.pasta_temporaria("leitura_docx_sdt")
    d = docx.Document()
    p = d.add_paragraph("Início. ")
    p._p.append(parse_xml(
        '<w:sdt xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f'<w:sdtContent><w:r><w:t>Fiador: {F.PESSOA_2}</w:t></w:r></w:sdtContent></w:sdt>'))
    d.save(str(pasta / "s.docx"))
    r = ler_docx(pasta / "s.docx")
    i = r.texto.index("Fiador: " + F.PESSOA_2)
    assert isinstance(r.origens[i], OrigemDocx)


def test_imagem_ou_objeto_embutido_gera_aviso():
    import io
    import zipfile
    import docx
    import fitz
    pasta = F.pasta_temporaria("leitura_docx_imagem")
    assert not any("imagem" in a for a in ler_docx(F.contrato_docx(pasta / "c.docx")).avisos)
    d = docx.Document(str(pasta / "c.docx"))
    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 8, 8), False)
    d.add_picture(io.BytesIO(pix.tobytes("png")))
    d.save(str(pasta / "img.docx"))
    aviso = "o documento tem imagem(ns) ou objeto(s) embutido(s) que não são lidos nem tratados; confira-os"
    assert aviso in ler_docx(pasta / "img.docx").avisos
    # objeto embutido (word/embeddings/...) sem imagem
    with zipfile.ZipFile(pasta / "c.docx") as z, zipfile.ZipFile(pasta / "obj.docx", "w") as w:
        for item in z.infolist():
            w.writestr(item, z.read(item.filename))
        w.writestr("word/embeddings/oleObject1.bin", b"\0" * 16)
    assert aviso in ler_docx(pasta / "obj.docx").avisos


if __name__ == "__main__":
    from _runner import rodar
    rodar(globals())
