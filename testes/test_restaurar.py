# SPDX-License-Identifier: AGPL-3.0-only
# Termo adicional (AGPL-3.0, seção 7(b)): preservar os avisos de autoria; ver CREDITOS.md.
import _base  # noqa: F401
import fabrica as F
from pseudonimizador.leitura import ler
from pseudonimizador.motor.trocas import aplicar_trocas
from pseudonimizador.restaurar import TokenOrfao, mapa_do_docx, restaurar_docx, restaurar_texto
from pseudonimizador.revisao import Revisao
from pseudonimizador.saida import gerar


def test_restaurar_docx_devolve_o_texto_original():
    pasta = F.pasta_temporaria("restaurar")
    doc = ler(F.contrato_docx(pasta / "contrato.docx"))
    g = gerar(doc, Revisao(doc).fechar(), formato="layout", modo="reversivel",
              pasta_saida=pasta / "Saida", pasta_mapas=pasta / "Mapas")
    mapa = mapa_do_docx(g.saida, pasta / "Mapas")
    assert mapa is not None
    destino = pasta / "restaurado.docx"
    assert restaurar_docx(g.saida, destino, mapa) > 0
    assert ler(destino).texto == doc.texto


def test_restaurar_texto_e_token_orfao():
    pasta = F.pasta_temporaria("restaurar2")
    doc = ler(F.contrato_docx(pasta / "contrato.docx"))
    res = Revisao(doc).fechar()
    protegido = aplicar_trocas(doc.texto, res.trocas)
    assert restaurar_texto(protegido, res.mapa) == doc.texto
    try:
        restaurar_texto("resposta citando [[PESSOA_99]]", res.mapa)
    except TokenOrfao as e:
        assert "[[PESSOA_99]]" in str(e)
    else:
        raise AssertionError("tóken órfão deveria ser apontado")


def test_restaurar_docx_alvo_de_hyperlink_e_token_orfao():
    import docx
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    pasta = F.pasta_temporaria("restaurar3")
    mapa = {"tokens": {"[[EMAIL_1]]": {"valor": F.EMAIL}}}
    d = docx.Document()
    rid = d.part.relate_to("mailto:[[EMAIL_1]]", RT.HYPERLINK, is_external=True)
    d.add_paragraph("oi")
    d.save(str(pasta / "a.docx"))
    assert restaurar_docx(pasta / "a.docx", pasta / "b.docx", mapa) == 1
    assert docx.Document(str(pasta / "b.docx")).part.rels[rid].target_ref == "mailto:" + F.EMAIL
    d.add_paragraph("cita [[PESSOA_9]]")
    d.save(str(pasta / "c.docx"))
    try:
        restaurar_docx(pasta / "c.docx", pasta / "d.docx", mapa)
    except TokenOrfao:
        assert not (pasta / "d.docx").exists()
    else:
        raise AssertionError("tóken órfão deveria ser apontado")


if __name__ == "__main__":
    from _runner import rodar
    rodar(globals())
