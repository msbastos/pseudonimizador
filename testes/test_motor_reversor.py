# SPDX-License-Identifier: AGPL-3.0-only
# Termo adicional (AGPL-3.0, seção 7(b)): preservar os avisos de autoria; ver CREDITOS.md.
r"""Testes: reversor de .docx — tokens [[...]] trocados pelos dados
reais do mapa, run-aware (preserva a formatacao de runs vizinhos),
com erro fatal para tóken orfao e gravacao atomica.

So' exercita a ESTRUTURA do mapa (dict com "tokens": {token: {"valor",
...}}) — nao importa pseudonimizador.motor.mapa, para que quem
gera o mapa e quem o reverte fiquem desacoplados.

Se python-docx nao estiver instalado, os testes se anulam sozinhos
(o reversor so' precisa da biblioteca para as funcoes que mexem em
.docx; despseudonimizar_texto nem a importa).

Rodar:  python -X utf8 testes/test_reversor.py
"""
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))  # raiz do projeto

try:
    import docx
except ImportError:
    print("python-docx ausente — testes do reversor pulados (rode INSTALAR.bat)")
    sys.exit(0)

MAPA = {"schema_version": 1, "cnj": "0", "tokens": {
    "[[RECLAMANTE_1]]": {"valor": "João da Silva", "categoria": "pessoa"},
    "[[CPF_1]]": {"valor": "529.982.247-25", "categoria": "cpf"}}}


def test_texto_simples():
    from pseudonimizador.motor.reversor import despseudonimizar_texto
    assert despseudonimizar_texto("autor [[RECLAMANTE_1]], CPF [[CPF_1]]", MAPA) \
        == "autor João da Silva, CPF 529.982.247-25"


def test_token_orfao_e_erro_fatal():
    from pseudonimizador.motor.reversor import despseudonimizar_texto, TokenOrfao
    try:
        despseudonimizar_texto("[[FANTASMA_9]]", MAPA)
        assert False, "deveria ter levantado TokenOrfao para tóken sem entrada no mapa"
    except TokenOrfao:
        pass


def test_docx_token_dividido_entre_runs_preserva_formatacao():
    from pseudonimizador.motor.reversor import despseudonimizar_docx
    with tempfile.TemporaryDirectory() as td:
        caminho = Path(td) / "minuta.docx"
        doc = docx.Document()
        par = doc.add_paragraph()
        run1 = par.add_run("réu de [[RECLA")
        run2 = par.add_run("MANTE_1]] em juízo")
        run2.bold = True
        tabela = doc.add_table(rows=1, cols=1)
        tabela.cell(0, 0).text = "[[CPF_1]]"
        doc.save(str(caminho))

        trocas = despseudonimizar_docx(caminho, MAPA)
        assert trocas == 2  # 1 no paragrafo (dividido em 2 runs) + 1 na tabela

        doc2 = docx.Document(str(caminho))
        par2 = doc2.paragraphs[0]
        assert "".join(r.text for r in par2.runs) == "réu de João da Silva em juízo"
        # run1 (nao-bold) continua nao-bold — recebeu o valor real na frente.
        assert par2.runs[0].bold is not True
        assert "João da Silva" in par2.runs[0].text
        # o run que era bold segue existindo, ainda bold, so' perdeu o
        # trecho do tóken — o resto ("em juízo") continua com a formatacao.
        bold_runs = [r for r in par2.runs if r.bold]
        assert bold_runs, "o run bold desapareceu"
        assert "em juízo" in "".join(r.text for r in bold_runs)

        assert doc2.tables[0].cell(0, 0).text == "529.982.247-25"

        # nenhum "[[" resta em paragrafo nenhum, dentro ou fora de tabela.
        for p in doc2.paragraphs:
            assert "[[" not in p.text
        for t in doc2.tables:
            for row in t.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        assert "[[" not in p.text


def test_docx_tem_token():
    from pseudonimizador.motor.reversor import docx_tem_token
    with tempfile.TemporaryDirectory() as td:
        com_token = Path(td) / "com_token.docx"
        doc = docx.Document()
        doc.add_paragraph("texto [[CPF_1]] aqui")
        doc.save(str(com_token))
        assert docx_tem_token(com_token) is True

        sem_token = Path(td) / "sem_token.docx"
        doc2 = docx.Document()
        doc2.add_paragraph("texto sem nada de especial")
        doc2.save(str(sem_token))
        assert docx_tem_token(sem_token) is False


# --- casos extras (autorrevisao do brief) ---------------------------------

def test_texto_com_token_repetido_e_no_fim_do_paragrafo():
    from pseudonimizador.motor.reversor import despseudonimizar_texto
    assert despseudonimizar_texto(
        "[[CPF_1]] e de novo [[CPF_1]], reclamante [[RECLAMANTE_1]]", MAPA) \
        == "529.982.247-25 e de novo 529.982.247-25, reclamante João da Silva"


def test_docx_mesmo_token_duas_vezes_no_mesmo_paragrafo():
    from pseudonimizador.motor.reversor import despseudonimizar_docx
    with tempfile.TemporaryDirectory() as td:
        caminho = Path(td) / "repetido.docx"
        doc = docx.Document()
        doc.add_paragraph("[[CPF_1]] confere com [[CPF_1]]")
        doc.save(str(caminho))

        trocas = despseudonimizar_docx(caminho, MAPA)
        assert trocas == 2

        doc2 = docx.Document(str(caminho))
        assert doc2.paragraphs[0].text == \
            "529.982.247-25 confere com 529.982.247-25"


def test_docx_token_no_cabecalho_e_rodape():
    # Regressao: _todos_paragrafos so' cobria doc.paragraphs e tabelas soltas
    # — um tóken em cabecalho/rodape sobrevivia a reversao, e docx_tem_token
    # dizia False mesmo com dado sensivel intacto no arquivo.
    from pseudonimizador.motor.reversor import (despseudonimizar_docx,
                                                       docx_tem_token)
    with tempfile.TemporaryDirectory() as td:
        caminho = Path(td) / "cabecalho_rodape.docx"
        doc = docx.Document()
        secao = doc.sections[0]
        secao.header.paragraphs[0].text = "Processo de [[RECLAMANTE_1]]"
        secao.footer.paragraphs[0].text = "CPF [[CPF_1]]"
        doc.add_paragraph("corpo normal, sem tóken")
        doc.save(str(caminho))

        assert docx_tem_token(caminho) is True

        trocas = despseudonimizar_docx(caminho, MAPA)
        assert trocas == 2

        doc2 = docx.Document(str(caminho))
        secao2 = doc2.sections[0]
        assert secao2.header.paragraphs[0].text == "Processo de João da Silva"
        assert secao2.footer.paragraphs[0].text == "CPF 529.982.247-25"
        assert docx_tem_token(caminho) is False


def test_docx_sem_nenhum_token_retorna_zero_e_conteudo_intocado():
    # despseudonimizar_docx sempre grava (atomico) mesmo sem tóken algum —
    # mais simples e previsivel do que condicionar o save ao numero de
    # trocas; o conteudo textual sai identico, so' a serializacao do
    # arquivo em disco pode diferir byte a byte.
    from pseudonimizador.motor.reversor import despseudonimizar_docx
    with tempfile.TemporaryDirectory() as td:
        caminho = Path(td) / "sem_token.docx"
        doc = docx.Document()
        doc.add_paragraph("nada para trocar aqui")
        doc.save(str(caminho))

        trocas = despseudonimizar_docx(caminho, MAPA)
        assert trocas == 0

        doc2 = docx.Document(str(caminho))
        assert doc2.paragraphs[0].text == "nada para trocar aqui"


if __name__ == "__main__":
    testes = [(n, f) for n, f in sorted(globals().items())
              if n.startswith("test_") and callable(f)]
    falhas = 0
    for nome, fn in testes:
        try:
            fn()
            print(f"ok    {nome}")
        except AssertionError as e:
            falhas += 1
            print(f"FALHA {nome}\n      {e}")
        except Exception as e:  # noqa: BLE001
            falhas += 1
            print(f"FALHA {nome}\n      {e!r}")
    print(f"\n{len(testes) - falhas}/{len(testes)} passaram, {falhas} falha(s)")
    sys.exit(1 if falhas else 0)
